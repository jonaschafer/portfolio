#!/usr/bin/env python3
"""
Export the Aleph brand audit into app/aleph/audit/data.json, plus optimized images.

The audit folder in Dropbox is the source of truth. Rerun this whenever the
sheet, findings or takeaways change:

    python3 scripts/aleph-audit-export.py              # data + images
    python3 scripts/aleph-audit-export.py --no-images  # data only

Reads from ALEPH_AUDIT_DIR (default: ~/Dropbox/ jon/work/clients/active/aleph/26_0910 Audit):
  aleph-brand-audit.xlsx          Audit tab: surfaces, What it is, Issue / Opportunity (the read)
  findings.md                     F-001 onward (the evidence)
  takeaways.md                    fuller readout and "By area"
  delivery/one-pager-draft.md     the overview copy
  competitor-captures-2026-09-12/headlines-2026-09-12.jsonl
Also reads scripts/aleph-audit-figma.json (Figma node ids for the surface map).

Writes:
  app/aleph/audit/data.json       imported by the server component only
  public/aleph/audit/img/*.webp   served behind the /aleph gate (middleware.js);
                                  <key>.webp for the page, <key>-lg.webp for the lightbox

Sanitizing for the Aleph audience happens here: em dashes go, the private
Slack DM becomes "Daniel, Aug 28" style labels, and working-file paths become
readable source labels. Hand-curated bits (the move per takeaway, image picks,
the ask ledger) live in the CURATION section below. Nothing is read from the
private field guide.
"""
import json
import os
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
from PIL import Image

REPO = Path(__file__).resolve().parent.parent
AUDIT = Path(os.environ.get(
    'ALEPH_AUDIT_DIR', '~/Dropbox/ jon/work/clients/active/aleph/26_0910 Audit')).expanduser()
OUT = REPO / 'app/aleph/audit/data.json'
IMG_DIR = REPO / 'public/aleph/audit/img'
IMG_URL = '/aleph/audit/img'
FIGMA_JSON = REPO / 'scripts/aleph-audit-figma.json'
FIGMA_FILE = 'https://www.figma.com/design/RRa1sgLX5WWJdOErGQobv6/Aleph-Audit'
SHEET_URL = 'https://docs.google.com/spreadsheets/d/1CC-KJ88cXnDylzzwFHrnJgbAAJ0PugPx/edit'

WARNINGS = []

# ---------------------------------------------------------------------------
# Voice and privacy
# ---------------------------------------------------------------------------

VOICE_FLAGS = ['genuinely', 'honestly', 'straightforward', 'leverage', 'robust',
               'unlock', 'holistic', 'synergy', 'best-in-class']


def clean(s):
    """Normalize a cell or paragraph: trim, drop em dashes, flag banned words."""
    if s is None:
        return None
    s = str(s).replace('\r', '').strip()
    if not s:
        return None
    s = re.sub(r'\s*—\s*', ', ', s)      # em dash
    s = re.sub(r'\s+–\s+', ', ', s)      # spaced en dash used as a dash
    s = re.sub(r'[ \t]+', ' ', s)
    for w in VOICE_FLAGS:
        if re.search(r'\b' + re.escape(w) + r'\b', s, re.I):
            WARNINGS.append(f'voice: "{w}" in "{s[:70]}..." (fine if quoted)')
    return s


PRIVATE = [r'\bcomp\b', r'\$2\d\dk', r'\bFMLA\b', r'surgery', r'two weeks\' notice', r'\bexit\b',
           r'offer letter', r'field guide', r'daniel-dm', r'gdm-', r'sources/slack', r'\bprivate\b']


def assert_public(obj, where='data'):
    """Fail the export if anything private slipped through."""
    text = json.dumps(obj, ensure_ascii=False)
    for pat in PRIVATE:
        for m in re.finditer(pat, text, re.I):
            ctx = text[max(0, m.start() - 50): m.end() + 50]
            # "exit" is fine as HubSpot's exit modal / exit-intent pop-up
            if pat == r'\bexit\b' and re.search(r'exit[- ](modal|intent|pop)', ctx, re.I):
                continue
            raise SystemExit(f'PRIVATE pattern {pat!r} in {where}: ...{ctx}...')
    if '—' in text:
        raise SystemExit(f'em dash left in {where}')


def source_chip(part):
    """One entry from the sheet's Source column -> {label, kind}."""
    p = clean(part)
    m = re.match(r'Slack, ([\w-]+), ([^,]+), (.+)$', p)
    if m:
        chan, who, when = m.groups()
        if chan == 'daniel-dm' or chan.startswith('gdm-'):
            return {'label': f'{who}, {when}', 'kind': 'slack'}
        return {'label': f'#{chan}, {who}, {when}', 'kind': 'slack'}
    m = re.match(r'Granola, (.+?) \((.+)\)$', p)
    if m:
        return {'label': f'{m.group(1)[0].upper()}{m.group(1)[1:]}, {m.group(2)}', 'kind': 'meeting'}
    m = re.match(r'(.+?) \((.+)\)$', p)
    if m:
        p = f'{m.group(1)}, {m.group(2)}'
    kind = 'jon' if p.startswith('Jon,') else 'doc'
    return {'label': p, 'kind': kind}


# Working-file paths in findings.md -> labels a reader can use. Output keeps the
# order each label first appears.
SOURCE_RULES = [
    (r'Slack, (?:daniel-dm|gdm-[\w-]+), ([A-Z][a-z]+), ([A-Z][a-z]{2} \d+(?: to \d+)?)', r'\1, Slack, \2'),
    (r'Slack, (?:daniel-dm|gdm-[\w-]+), ([A-Z][a-z]{2} \d+(?: to \d+)?)', r'Slack, \1'),
    (r'daniel-dm', 'Daniel, Slack'),
    (r'brand-revamp/brand-revamp', '#brand-revamp'),
    (r'aleph-red-antler-kickoff', 'Red Antler kickoff notes, Aug 26'),
    (r'granola-ko', 'Ads kickoff notes, Jul 30'),
    (r'brand-survey', 'Internal brand survey'),
    (r'g2-reviews', 'G2 reviews'),
    (r'reddit-ai-sentiment', 'Reddit, AI threads'),
    (r'reddit-mentions', 'Reddit, Aleph mentions'),
    (r'linkedin-x-web', 'LinkedIn, X and web search'),
    (r'sources/culture/\*', 'G2, Reddit, LinkedIn and X search'),
    (r'aleph-namesakes', 'Tech and search notes, Sep 10'),
    (r'serp-', 'Google results, Sep 10'),
    (r'sw-\*|similarweb', 'Similarweb'),
    (r'gtm\.js', 'Tag manager container'),
    (r'aleph\.html|Webflow shared CSS', 'Site source and CSS'),
    (r'sitemap', 'Sitemap'),
    (r'whois', 'Domain lookup'),
    (r'TEmp|Tech stack', 'Figma, tech stack'),
    (r'cta-inventory', 'CTA inventory'),
    (r'flow-map', 'Conversion flow map'),
    (r'recon/(?!fitness)', 'Form and page recon'),
    (r'variants', 'Variant test log'),
    (r'fitness-explore|fpa-fitness-test|fitness app bundle|app bundle', 'Fitness test app'),
    (r'captures? \d|capture \d', 'Conversion captures'),
    (r'Aleph Audit > Conversion', 'Figma, conversion captures'),
    (r'competitor-captures', 'Competitor homepages, Sep 12'),
    (r'identity-usage', 'Identity usage captures, Sep 13'),
    (r'listing pages', 'Marketplace listings'),
    (r'app shell', 'App shell, Sep 13'),
    (r'excel-addin', 'Excel add-in site'),
    (r'Jon, Sep 11', 'Jon, Sep 11'),
]


def finding_sources(raw):
    hits = []
    for pat, label in SOURCE_RULES:
        for m in re.finditer(pat, raw):
            lab = m.expand(label) if '\\' in label else label
            hits.append((m.start(), lab))
    seen, out = set(), []
    for _, lab in sorted(hits):
        if lab not in seen:
            seen.add(lab)
            out.append(lab)
    if not out:
        WARNINGS.append(f'source: nothing readable in "{raw[:60]}"')
    return out


# ---------------------------------------------------------------------------
# CURATION: hand-picked, voice-checked. Keyed by the lead F-number of each
# one-pager item so reordering the list doesn't break anything.
# ---------------------------------------------------------------------------

# One-pager item (by lead F-number) -> takeaways.md item number (fuller read)
FULLER = {'F-008': 1, 'F-014': 2, 'F-017': 4, 'F-015': 5, 'F-002': 6, 'F-001': 7}
# takeaways.md items that ride along with a one-pager item as "related": (item number, its F-refs)
RELATED = {'F-014': (3, ['F-029'])}

MOVES = {
    'F-008': 'Give the customer voice a home next to the trust message, with real customers and the CX team in it.',
    'F-014': 'Build the rebrand\'s message around trust, and bring the proof: governed data and answers someone can check.',
    'F-017': 'One look and voice from the ad to the booked meeting, including forms, scheduler, emails and login. The Agentic Budgeting demo page is the model.',
    'F-015': 'Lead with the job it takes off someone\'s plate and how they check the result.',
    'F-002': 'Hand Red Antler name ownability as an input to the positioning line and the aleph.ai switch, and buy our brand terms now.',
    'F-001': 'Set a performance budget before the rebrand adds more motion.',
}

VISUALS = {
    'F-008': {'type': 'quotes', 'stat': {'value': '4.9', 'label': 'on G2, from 108 reviews'}, 'quotes': [
        {'text': 'speed of integration and cost makes it a no brainer over Vena', 'by': 'Reddit, r/FPandA'},
        {'text': 'used it for over 3 years now and it\'s been awesome', 'by': 'Reddit, r/FPandA'},
    ]},
    'F-014': {'type': 'quotes', 'quotes': [
        {'text': 'AI can guess. Finance still has to sign its name to the number.', 'by': 'A Vena partner, LinkedIn'},
        {'text': 'It\'s only as good as your data pipelines.', 'by': 'r/FPandA, on the ChatGPT for Financial Services launch'},
        {'text': 'answers I actually trust', 'by': 'An Aleph customer on G2, querying Aleph from Claude'},
    ]},
    'F-017': {'type': 'strip', 'images': ['flow-01', 'flow-02', 'flow-03', 'flow-04', 'flow-05', 'flow-06', 'flow-07']},
    'F-015': {'type': 'quotes', 'quotes': [
        {'text': 'Does anyone else\'s boss have terminal AI brain?', 'by': 'Top r/FPandA thread on agents, 120 upvotes'},
        {'text': 'Nope but our executives will still pretend we\'re an \'AI first\' company.', 'by': 'Top reply to "Success story for using AI/agent?"'},
    ]},
    'F-002': {'type': 'image', 'image': 'serp-brand'},
    'F-001': {'type': 'stat', 'stat': {'value': '32.2 s', 'label': 'to show the main content on a phone'},
              'support': [{'value': '2.5 s', 'label': 'Google\'s target'}, {'value': '8.6 MB', 'label': 'page weight'},
                          {'value': '26', 'label': 'Lighthouse mobile score, of 100'}]},
}

# Images: key -> (source path in the audit folder, crop box or None, page width, alt, caption)
IMAGES = {
    'flow-01': ('conversion-flow/captures/01-entry-banner.png', None, 724, 'getaleph.com with the HubSpot top banner', 'The site, with HubSpot\'s top banner'),
    'flow-02': ('conversion-flow/captures/02-entry-exit-modal.png', None, 724, 'HubSpot exit pop-up offering a 15-minute call', 'HubSpot exit pop-up: 15 minutes'),
    'flow-03': ('conversion-flow/captures/03-demo-landing.png', (0, 0, 724, 450), 724, 'The /demo page, its own shell with no nav', '/demo: its own shell, 30 minutes'),
    'flow-04': ('conversion-flow/captures/10-trial-form.png', None, 724, 'HubSpot embedded form on the trial page', 'HubSpot form styling'),
    'flow-05': ('conversion-flow/captures/07-abdemo-landing.png', None, 724, 'Agentic Budgeting demo page', 'Agentic Budgeting demo: the model'),
    'flow-06': ('conversion-flow/captures/18-fitness-gate.png', (0, 0, 724, 450), 724, 'FP&A fitness test contact gate on Replit', 'Fitness test on Replit, gated'),
    'flow-07': ('conversion-flow/captures/19-signin-login.png', None, 724, 'Stock Auth0 sign-in page with a blue button', 'Auth0 sign-in, stock'),
    'serp-brand': ('tech-and-search-2026-09-10/figma-crops/serp-brand-1.png', (0, 0, 1920, 1560), 1100, 'Google results for "aleph": AI answer, Wikipedia, then getaleph.com', 'Google, "aleph," logged out, Sep 10. We\'re the third result.'),
    'serp-ai': ('tech-and-search-2026-09-10/figma-crops/serp-ai-1.png', (0, 0, 1920, 1560), 1100, 'Google results for "ai fp&a software" with Aleph first', 'Google, "ai fp&a software," Sep 10. We\'re the top organic result.'),
    'sw-channels': ('tech-and-search-2026-09-10/figma-crops/sw-channels.png', None, 1100, 'Similarweb channel mix for getaleph.com', 'Similarweb channel mix, Jun to Aug (free estimates)'),
    'linkedin': ('social-captures-2026-09-11/01-linkedin.png', (0, 0, 724, 900), 724, 'Aleph LinkedIn company page, logged out', 'Our LinkedIn page, logged out, Sep 11'),
    'trial': ('conversion-flow/captures/09-trial-landing.png', (0, 0, 724, 560), 724, 'Trial page promising a free, full access trial', '/trial promises full access in minutes'),
    'contact': ('conversion-flow/captures/13-contact-landing.png', (0, 0, 724, 560), 724, 'Contact page form', '/contact-us: six required fields'),
    'home-aleph': ('voice-examples-2026-09-13/pos-01-home.png', None, 900, 'Aleph homepage first screen', 'Aleph: "AI-native FP&A that connects your data, strategy, and spreadsheets"'),
    'home-drivetrain': ('competitor-captures-2026-09-12/03-drivetrain.png', (0, 0, 724, 500), 724, 'Drivetrain homepage first screen', 'Drivetrain: "The AI-native FP&A platform that does the work"'),
    'icon-excel': ('identity-usage-2026-09-13/microsoft-excel-icon.png', None, 200, 'Aleph symbol in blue on the Microsoft Marketplace', 'Microsoft Marketplace'),
    'icon-google': ('identity-usage-2026-09-13/google-workspace-icon.png', None, 200, 'Aleph symbol in black on Google Workspace Marketplace', 'Google Workspace'),
}
LIGHTBOX_WIDTH = 1800   # cap for the <key>-lg.webp lightbox variant

# Evidence thumbnails per finding
FINDING_IMAGES = {
    'F-002': ['serp-brand'], 'F-005': ['serp-ai'], 'F-006': ['sw-channels'], 'F-011': ['linkedin'],
    'F-017': ['trial'], 'F-018': ['flow-06'], 'F-019': ['flow-01', 'flow-02', 'flow-03', 'flow-04', 'flow-05', 'flow-06', 'flow-07'],
    'F-020': ['flow-02'], 'F-021': ['flow-03'], 'F-023': ['flow-05'], 'F-025': ['contact'], 'F-026': ['flow-07'],
    'F-029': ['home-aleph', 'home-drivetrain'], 'F-030': ['icon-excel', 'icon-google'],
}

# Homepage headlines for the comparison slide (from headlines-2026-09-12.jsonl)
COMPARE_FINANCE = ['Aleph', 'Drivetrain', 'Campfire', 'Datarails', 'Concourse', 'Anaplan', 'Runway (cfo.ai)']
COMPARE_REFERENCE = ['Harvey', 'Ramp', 'Linear', 'Stripe', 'Anthropic']

# What Daniel (and Albert) asked for, in order. Private DM sources shown as the person.
LEDGER = [
    ('Jun 24', '#brand-revamp, Daniel', 'Agency applications to ship right away: five tentpole webpages (homepage, a webinar page, two or three core product pages), a new sales deck plus a piece or two of sales collateral, viral merch, and explorations for webinars and in-person events.'),
    ('Jul 30', 'Ads kickoff, Daniel', 'Ads and social profiles from scratch, ad templates, a giveaway ad, animation, and a design.md for AI generation.'),
    ('Aug 3', 'Red Antler brief', 'Red Antler\'s deliverables: symbol and wordmark (2D, 3D, animated, and app icons for Slack, Sheets and Excel), palette, type, voice guide, motion language, a visual style for diagrams and data, and applications (paid ads, social, blog heroes, ebook covers, apparel, the Zoom loop, the homepage). WildFruit does the 12-slide sales deck.'),
    ('Aug 7', 'Daniel', 'The Agentic Budgeting launch sprint: video, landing page, ebook, launch assets.'),
    ('Aug 27', '#brand-design, Albert', 'The ChatGPT plugin page is text only, and Hex\'s has product visuals. Next day Daniel asked for MCP page visuals and a new MCP and Agent icon.'),
    ('Aug 28', 'Daniel', 'The core ask: map "the \'not in the official deliverable list for RA but critical for us to execute\' stuff like the MCP screenshots."'),
    ('Sep 3', 'Daniel', 'Mac-style app icons for Claude, ChatGPT, Cursor and Aleph in the launch video, plus the logo animation as a standalone element for other videos.'),
    ('Sep 10', 'Daniel', 'Further-out investments like great merch "feel super related to the brand application exploration work."'),
    ('Sep 11', 'Daniel', '"A good amount of white space to add and propose surfaces and experiences." Named: a friendly video guide channel (like Framer\'s), an onboarding drip campaign ("We still need to educate users and guide them. That\'s non-negotiable."), and a public changelog and docs.'),
]

# Phrases the page shows with a highlighter mark (==...== in the data)
MARKS = {
    'tldr': ['AI they can trust because the data under it is right'],
    'question': ['customers who brought Aleph to a new company'],
}


def mark(text, key):
    for phrase in MARKS.get(key, []):
        if text and phrase in text:
            text = text.replace(phrase, f'=={phrase}==', 1)
        else:
            WARNINGS.append(f'mark: "{phrase}" not found in {key}')
    return text


# ---------------------------------------------------------------------------
# Parsers
# ---------------------------------------------------------------------------

FREF = re.compile(r'F-\d{3}')


def refs_in(s):
    return sorted(set(FREF.findall(s or '')))


def parse_one_pager():
    text = (AUDIT / 'delivery/one-pager-draft.md').read_text(encoding='utf-8')
    body = text.split('\n# Aleph brand audit\n', 1)[1].split('\n## Build notes', 1)[0]
    body = body.rsplit('\n---', 1)[0]
    sections = re.split(r'\n## ', body)
    head, rest = sections[0], {s.split('\n', 1)[0].strip(): s.split('\n', 1)[1].strip() for s in sections[1:]}
    paras = [p.strip() for p in head.split('\n\n') if p.strip()]
    out = {}
    for p in paras:
        if p.startswith('Jon ·'):
            continue
        if p.startswith('**Why an audit'):
            out['why'] = clean(re.sub(r'^\*\*Why an audit, and why now\.\*\*\s*', '', p))
        elif p.startswith('Quick flag'):
            out['bias'] = clean(p)
        elif p.startswith('**TL;DR:**'):
            out['tldr'] = clean(p.replace('**TL;DR:**', ''))
        else:
            out['intro'] = clean(p)

    items = []
    for line in rest['What stands out'].split('\n'):
        m = re.match(r'^(\d+)\. \*\*(.+?)\*\*\s*(.*?)\s*\[(F-[\d, F-]+)\]\s*$', line.strip())
        if m:
            refs = FREF.findall(m.group(4))
            items.append({'n': int(m.group(1)), 'claim': clean(m.group(2)), 'body': clean(m.group(3)), 'refs': refs})
    working = rest["What's working"]
    out['working'] = {'text': clean(re.sub(r'\s*\[F-[\d, F-]+\]\s*$', '', working)), 'refs': FREF.findall(working)}
    out['rebrand'] = [clean(l[2:]) for l in rest['What it means for the rebrand'].split('\n') if l.startswith('- ')]
    out['stillToDo'] = clean(rest['Still to do'])
    ask_paras = [p.strip() for p in rest['One ask'].split('\n\n') if p.strip()]
    # The doc's ask has a date placeholder and "comments right here"; the page keeps the substance.
    ask = re.sub(r'^By \[[^\]]+\], ', '', ask_paras[0])
    ask = ask.replace(' Comments right here are perfect.', '').replace('will fold them into', 'will fold your notes into')
    out['ask'] = clean(ask[0].upper() + ask[1:])
    out['askClose'] = clean(ask_paras[1]) if len(ask_paras) > 1 else None
    out['standouts'] = items
    return out


def parse_takeaways():
    text = (AUDIT / 'takeaways.md').read_text(encoding='utf-8')
    secs = {s.split('\n', 1)[0].strip(): s.split('\n', 1)[1] for s in re.split(r'\n## ', text)[1:]}
    items = {}
    for m in re.finditer(r'^(\d+)\. \*\*(.+?)\*\*\s*(.+?)(?=\n\n\d+\. |\n*\Z)', secs['What stands out so far'], re.S | re.M):
        items[int(m.group(1))] = {'claim': clean(m.group(2)), 'text': clean(m.group(3).strip())}
    rebrand = [clean(p) for p in secs['What this means for the rebrand'].strip().split('\n\n') if p.strip()]
    areas = []
    area_text = secs['By area']
    for block in re.split(r'\n(?=\*\*)', area_text.strip()):
        m = re.match(r'\*\*(.+?)\*\*\s*\n(.*)', block, re.S)
        if not m or m.group(1).startswith('Question'):
            continue
        a = {'bucket': m.group(1)}
        for l in m.group(2).split('\n'):
            l = l.strip()
            for key, pre in (('notWorking', '- Not working: '), ('working', '- Working: '), ('oneThing', '- The one thing: ')):
                if l.startswith(pre):
                    a[key] = clean(l[len(pre):])
        areas.append(a)
    qm = re.search(r'\*\*Question:\*\*\s*(.+)', area_text)
    return {'items': items, 'rebrand': rebrand, 'areas': areas, 'question': clean(qm.group(1)) if qm else None}


def parse_findings():
    text = (AUDIT / 'findings.md').read_text(encoding='utf-8')
    wb = openpyxl.load_workbook(AUDIT / 'aleph-brand-audit.xlsx', read_only=True)
    sheet_rows = {r[0]: r for r in wb['Findings'].iter_rows(min_row=2, values_only=True) if r and r[0]}
    blocks = re.split(r'\n### ', text.split('\n---\n', 1)[1].split('\n## Changelog', 1)[0])
    out = []
    for b in blocks:
        b = b.strip().rstrip('-').strip()
        m = re.match(r'(F-\d{3}) · (.+)', b)
        if not m:
            continue
        fid, title = m.groups()
        fields, cur = {}, None
        for line in b.split('\n')[1:]:
            fm = re.match(r'- \*\*(.+?):\*\*\s*(.*)', line)
            if fm:
                cur = fm.group(1)
                fields[cur] = fm.group(2)
            elif cur and line.strip():
                fields[cur] += '\n' + line.strip()
        links, figma = [], None
        for lm in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)', fields.get('Links', ''), re.S):
            label, urls = lm.group(1), lm.group(2).split()
            for i, u in enumerate(urls):
                lab = label if i == 0 else re.sub(r'^https?://(www\.)?([^/]+).*', r'\2', u)
                if 'figma.com' in u:
                    figma = figma or {'label': clean(label.replace('Figma: ', '')), 'url': u}
                else:
                    links.append({'label': clean(lab), 'url': u})
        conf = fields.get('Confidence', '')
        levels = []
        for key, pat in (('hard', r'hard (data|survey)'), ('directional', r'directional'), ('hunch', r'hunch')):
            pos = re.search(pat, conf, re.I)
            if pos:
                levels.append((pos.start(), key))
        levels = [k for _, k in sorted(levels)] or ['directional']
        bs = fields.get('Bucket / surface', '')
        bucket, _, surface = bs.partition(' / ')
        row = sheet_rows.get(fid)
        out.append({
            'id': fid,
            'slug': fid.lower(),
            'claim': clean(title),
            'bucket': clean(row[2]) if row else clean(bucket),
            'surface': clean(surface),
            'surfaceRaw': bs + ' ; ' + (row[3] if row and row[3] else ''),
            'evidence': clean(fields.get('Evidence', '')),
            'confidence': clean(conf),
            'levels': levels,
            'sources': finding_sources(fields.get('Source', '')),
            'links': links,
            'figma': figma,
            'date': fields.get('Date', '').strip(),
            'images': FINDING_IMAGES.get(fid, []),
        })
    return out


def figma_link(location, figma_map):
    """'IDENTITY, Logo; marketing, Video' -> first resolvable node link."""
    if not location or not figma_map:
        return None
    aliases = figma_map.get('aliases', {})
    heads = figma_map.get('headings', {})

    def strip(s):
        return re.sub(r'\s*\([^)]*\)\s*$', '', s).strip()

    for part in location.lower().split(';'):
        part = part.strip()
        segs = [x.strip() for x in part.split(',')]
        # most specific first: whole part, then the sub label, then the section
        for key in (part, strip(part), segs[-1], strip(segs[-1]), segs[0], strip(segs[0])):
            k = aliases.get(key, key)
            if k == '__none__':          # a placeholder card, nothing to link to
                return None
            if k in heads:
                return f'{FIGMA_FILE}?node-id={heads[k].replace(":", "-")}'
    WARNINGS.append(f'figma: no node for location "{location}"')
    return None


def parse_surfaces(findings, figma_map):
    wb = openpyxl.load_workbook(AUDIT / 'aleph-brand-audit.xlsx', read_only=True)
    rows = list(wb['Audit'].iter_rows(values_only=True))
    col = {h: i for i, h in enumerate(rows[0]) if h}
    for need in ('Surface', 'What it is', 'Issue (1 line)', 'Opportunity (1 line)', 'In RA SOW?', 'Priority'):
        if need not in col:
            raise SystemExit(f'Audit tab is missing the "{need}" column')

    def g(r, name):
        i = col.get(name)
        return clean(r[i]) if i is not None and i < len(r) else None

    surfaces = []
    for r in rows[1:]:
        if not r or not r[col['Surface']]:
            continue
        raw_name = r[col['Surface']]
        what = g(r, 'What it is')
        if not what:
            WARNINGS.append(f'surface: "{raw_name}" has no "What it is" on the Audit tab')
        # Findings: F-numbers in the read, plus any finding whose surface names this one
        fr = set(refs_in(' '.join(filter(None, [g(r, 'Issue (1 line)'), g(r, 'Opportunity (1 line)')]))))
        for f in findings:
            if raw_name in f['surfaceRaw']:
                fr.add(f['id'])
        in_fig = g(r, 'In Figma?')
        surfaces.append({
            'group': g(r, 'Group'), 'sub': g(r, 'Sub'), 'name': clean(raw_name),
            'what': what,
            'issue': g(r, 'Issue (1 line)'), 'opportunity': g(r, 'Opportunity (1 line)'),
            'sow': g(r, 'In RA SOW?'), 'priority': g(r, 'Priority'),
            'state': g(r, 'State'), 'pov': g(r, 'POV'), 'craft': g(r, 'Craft'),
            'owner': g(r, 'Owner'),
            'sources': [source_chip(p) for p in (g(r, 'Source') or '').split(';') if p.strip()],
            'inFigma': in_fig,
            'figmaUrl': figma_link(g(r, 'Figma location'), figma_map) if in_fig and in_fig != 'No' else None,
            'findings': sorted(fr),
        })
    return surfaces


def rollup(surfaces):
    def count(key, val):
        return sum(1 for s in surfaces if s[key] == val)
    ours = [s for s in surfaces if s['sow'] != 'Yes']
    ours_p1 = [s for s in ours if s['priority'] == 'P1']
    return {
        'total': len(surfaces),
        'sow': {k: count('sow', k) for k in ('Yes', 'Partial', 'No')},
        'ours': len(ours),
        'priority': {k: count('priority', k) for k in ('P1', 'P2', 'P3')},
        'oursP1': [s['name'] for s in ours_p1],
    }


def headlines():
    rows = {}
    path = AUDIT / 'competitor-captures-2026-09-12/headlines-2026-09-12.jsonl'
    for line in path.read_text(encoding='utf-8').splitlines():
        d = json.loads(line)
        h1 = d['h1'][0]
        half = len(h1) // 2
        if len(h1) % 2 == 1 and h1[:half] == h1[half + 1:]:   # "X X" doubled by the crawler
            h1 = h1[:half]
        h1 = re.split(r'(?<=\.) ', h1)[0] if len(h1) > 90 else h1
        rows[d['site']] = {'site': d['site'], 'h1': clean(h1), 'title': clean(d['title']),
                           'aiInTitle': bool(re.search(r'\bAI\b', d['title']))}
    return {'finance': [rows[s] for s in COMPARE_FINANCE if s in rows],
            'reference': [rows[s] for s in COMPARE_REFERENCE if s in rows]}


def export_images(write=True):
    meta = {}
    if write:
        IMG_DIR.mkdir(parents=True, exist_ok=True)
    for key, (src, crop, width, alt, caption) in IMAGES.items():
        out, out_lg = IMG_DIR / f'{key}.webp', IMG_DIR / f'{key}-lg.webp'
        if write:
            im = Image.open(AUDIT / src).convert('RGB')
            if crop:
                im = im.crop(crop)
            lg = im if im.width <= LIGHTBOX_WIDTH else im.resize((LIGHTBOX_WIDTH, round(im.height * LIGHTBOX_WIDTH / im.width)), Image.LANCZOS)
            lg.save(out_lg, 'WEBP', quality=80, method=6)
            if im.width > width:
                im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
            im.save(out, 'WEBP', quality=74, method=6)
        with Image.open(out) as done:
            w, h = done.size
        meta[key] = {'src': f'{IMG_URL}/{key}.webp', 'full': f'{IMG_URL}/{key}-lg.webp',
                     'w': w, 'h': h, 'alt': alt, 'caption': clean(caption)}
    return meta


def main():
    write_images = '--no-images' not in sys.argv
    figma_map = json.loads(FIGMA_JSON.read_text()) if FIGMA_JSON.exists() else None
    if not figma_map:
        WARNINGS.append('figma: scripts/aleph-audit-figma.json missing, surfaces get no Figma links')

    one = parse_one_pager()
    tk = parse_takeaways()
    findings = parse_findings()
    surfaces = parse_surfaces(findings, figma_map)
    images = export_images(write_images)

    standouts = []
    for it in one['standouts']:
        key = it['refs'][0]
        full = tk['items'].get(FULLER.get(key))
        rel = None
        if key in RELATED:
            num, rel_refs = RELATED[key]
            if num in tk['items']:
                rel = {**tk['items'][num], 'refs': rel_refs, 'images': ['home-aleph', 'home-drivetrain']}
        if key not in MOVES:
            WARNINGS.append(f'takeaway {key}: no move in MOVES')
        standouts.append({**it, 'key': key.lower(), 'move': MOVES.get(key),
                          'fuller': full['text'] if full else None,
                          'related': rel, 'visual': VISUALS.get(key)})

    known = {f['id'] for f in findings}
    for ref in set(FREF.findall(json.dumps([one, surfaces, tk]))) - known:
        WARNINGS.append(f'ref {ref} is cited but not in findings.md')

    data = {
        'meta': {
            'title': 'Aleph brand audit', 'author': 'Jon Schafer', 'date': 'Sep 2026',
            'exported': date.today().isoformat(), 'figmaFile': FIGMA_FILE, 'sheet': SHEET_URL,
            'loom': os.environ.get('ALEPH_LOOM_URL') or None,
            'counts': {'findings': len(findings), 'surfaces': len(surfaces)},
        },
        'overview': {
            'intro': one.get('intro'), 'why': one.get('why'), 'bias': one.get('bias'),
            'tldr': mark(one.get('tldr'), 'tldr'),
            'standouts': standouts, 'working': one['working'], 'rebrand': one['rebrand'],
            'stillToDo': one['stillToDo'], 'question': mark(tk['question'], 'question'),
            'ask': one['ask'], 'askClose': one['askClose'],
        },
        'rebrandFull': tk['rebrand'],
        'byArea': tk['areas'],
        'surfaces': {
            'ask': 'the "not in the official deliverable list for RA but critical for us to execute" stuff like the MCP screenshots.',
            'askBy': 'Daniel, Aug 28',
            'ledger': [{'date': d, 'who': w, 'text': clean(t)} for d, w, t in LEDGER],
            'rollup': rollup(surfaces), 'rows': surfaces,
        },
        'findings': findings,
        'compare': headlines(),
        'images': images,
    }
    for f in data['findings']:
        f.pop('surfaceRaw', None)

    assert_public(data)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    r = data['surfaces']['rollup']
    size = sum((IMG_DIR / f'{k}.webp').stat().st_size for k in images) // 1024
    print(f'wrote {OUT.relative_to(REPO)}: {len(findings)} findings, {r["total"]} surfaces, '
          f'{len(standouts)} takeaways, {len(images)} images ({size} KB on the page, lightbox files load on click)')
    print(f'surfaces with a Figma link: {sum(1 for s in surfaces if s["figmaUrl"])} of {r["total"]}; '
          f'with a read: {sum(1 for s in surfaces if s["issue"] or s["opportunity"])}')
    for w in WARNINGS:
        print('  !', w)


if __name__ == '__main__':
    main()
